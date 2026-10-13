"""
LaborLens – Document Ingestion Service
Reads PDF/DOCX files, splits into chunks, embeds, stores in ChromaDB.
"""

import asyncio
import csv
import contextlib
import math
import os
import re
import shutil
import subprocess
import tempfile
import threading
import uuid
from io import BytesIO, StringIO
from pathlib import Path
from typing import List, Tuple, Dict, Optional

from pypdf import PdfReader
from pypdf.generic import ContentStream
from docx import Document as DocxDocument

try:
    import pytesseract
except ImportError:  # Optional OCR dependency.
    pytesseract = None

try:
    from pdf2image import convert_from_path
except ImportError:  # Optional scanned-PDF dependency.
    convert_from_path = None

try:
    from PIL import Image, ImageChops
except ImportError:  # Optional image dependency.
    Image = None
    ImageChops = None

try:
    import pypdfium2 as pdfium
except ImportError:  # Optional Poppler-free PDF renderer.
    pdfium = None

from routers.services.chroma_service import add_documents, embed_texts
from routers.services.ocr_checkpoint import OCRCheckpoint
from routers.services.document_text_cleaner import strip_document_footnotes
from routers.services.kb_categories import DEFAULT_KB_CATEGORY, normalize_kb_category
from routers.services.rag_policy import infer_document_year, prepare_chunks_for_indexing
from config import settings
import logging

logger = logging.getLogger(__name__)

INGESTION_STATUS = {
    "is_processing": False,
    "files_completed": 0,
    "files_failed": 0,
    "total_files": 0,
    "active_jobs": 0,
    "current_files": [],
    "last_error": None,
}
_INGESTION_STATUS_LOCK = threading.Lock()

CHUNK_SIZE = 1200
OCR_DPI = 220
_OCR_PAGE_LOCK = threading.Lock()
# Ignore small decorative images; inspect meaningful raster areas without
# decoding all embedded images into memory. OCR still runs one page at a time.
_PDF_IMAGE_MIN_PAGE_AREA = 0.04
_PDF_IMAGE_MIN_PIXELS = 12000
_PDF_IMAGE_MAX_FORM_DEPTH = 12
_PDF_IMAGE_MAX_REGIONS = 128
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
ARTICLE_RE = re.compile(
    r"\b(?:ART\.|Article)\s*(?P<number>\d+)\.\s*(?:\[[^\]]+\]\s*)?"
    r"(?P<title>[^.\n\r-]{2,120})",
    re.IGNORECASE,
)
STRUCTURAL_BOUNDARY_RE = re.compile(
    r"\b(?:Book|Title|Chapter)\s+(?:[IVXLCDM]+|\d+|One|Two|Three|Four|Five|Six|Seven)"
    r"\s*(?:[–—-]|�)\s*",
    re.IGNORECASE,
)


def _start_ingestion_status(source_name: str) -> None:
    """Register an ingestion job without exposing mutable shared state."""

    display_name = source_name or "document"
    with _INGESTION_STATUS_LOCK:
        if INGESTION_STATUS["active_jobs"] == 0:
            INGESTION_STATUS.update(
                files_completed=0,
                files_failed=0,
                total_files=0,
                current_files=[],
                last_error=None,
            )
        INGESTION_STATUS["active_jobs"] += 1
        INGESTION_STATUS["total_files"] += 1
        INGESTION_STATUS["is_processing"] = True
        INGESTION_STATUS["current_files"].append(display_name)


def _finish_ingestion_status(
    source_name: str,
    *,
    succeeded: bool,
    error: Optional[str] = None,
) -> None:
    display_name = source_name or "document"
    with _INGESTION_STATUS_LOCK:
        INGESTION_STATUS["files_completed"] += 1
        if not succeeded:
            INGESTION_STATUS["files_failed"] += 1
        if error:
            INGESTION_STATUS["last_error"] = error[:500]
        INGESTION_STATUS["active_jobs"] = max(
            0, INGESTION_STATUS["active_jobs"] - 1
        )
        with contextlib.suppress(ValueError):
            INGESTION_STATUS["current_files"].remove(display_name)
        INGESTION_STATUS["is_processing"] = INGESTION_STATUS["active_jobs"] > 0


def get_ingestion_status() -> Dict[str, object]:
    """Return a snapshot suitable for the admin polling endpoint."""

    with _INGESTION_STATUS_LOCK:
        snapshot = dict(INGESTION_STATUS)
        snapshot["current_files"] = list(INGESTION_STATUS["current_files"])
    snapshot["current_file"] = (
        snapshot["current_files"][0] if snapshot["current_files"] else None
    )
    return snapshot


def pdf_reader_from_file(filepath: str) -> PdfReader:
    return PdfReader(BytesIO(Path(filepath).read_bytes()))


def clean_text(text: str) -> str:
    text = strip_document_footnotes(text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


def clean_chunk_text(text: str) -> str:
    """Normalize extracted text while preserving blank-line paragraph boundaries."""

    text = strip_document_footnotes(text or "")
    paragraphs = []
    for raw_paragraph in re.split(r"(?:\r?\n\s*){2,}", text):
        lines = [re.sub(r"\s+", " ", line).strip() for line in raw_paragraph.splitlines()]
        paragraph = " ".join(line for line in lines if line)
        if paragraph:
            paragraphs.append(paragraph)
    return "\n\n".join(paragraphs)


_LEGAL_ABBREVIATIONS = re.compile(
    r"\b(?:ART|SEC|NO|S|RA|PD|BP|DO|GR|AM|D\.O|R\.A|P\.D|B\.P)\.",
    re.IGNORECASE,
)
_SENTENCE_BOUNDARY = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9(\u201c\"'])")


def _paragraph_ranges(text: str, start: int, end: int) -> List[Tuple[int, int]]:
    """Return non-empty paragraph spans inside a document/article span."""

    segment = text[start:end]
    return [
        (start + match.start(), start + match.end())
        for match in re.finditer(r"\S.*?(?=\n\n|\Z)", segment, flags=re.DOTALL)
    ]


def _sentence_ranges(text: str, start: int, end: int) -> List[Tuple[int, int]]:
    """Split an oversized paragraph at likely sentence endings."""

    paragraph = text[start:end]
    protected = _LEGAL_ABBREVIATIONS.sub(
        lambda match: match.group(0).replace(".", "\uE000"), paragraph
    )
    boundaries = [match.end() for match in _SENTENCE_BOUNDARY.finditer(protected)]
    if not boundaries:
        return [(start, end)]

    ranges = []
    cursor = 0
    for boundary in boundaries:
        while boundary < len(paragraph) and paragraph[boundary].isspace():
            boundary += 1
        if boundary > cursor:
            ranges.append((start + cursor, start + boundary))
            cursor = boundary
    if cursor < len(paragraph):
        ranges.append((start + cursor, end))
    return ranges


def _bounded_unit_ranges(text: str, start: int, end: int) -> List[Tuple[int, int]]:
    """Keep a paragraph whole when possible; otherwise split at sentences/words."""

    if end - start <= CHUNK_SIZE:
        return [(start, end)]

    bounded = []
    for sentence_start, sentence_end in _sentence_ranges(text, start, end):
        cursor = sentence_start
        while sentence_end - cursor > CHUNK_SIZE:
            limit = cursor + CHUNK_SIZE
            # Prefer punctuation or whitespace near the target. A single
            # unbroken OCR token may require a hard split to bound the embed input.
            split_at = max(text.rfind(mark, cursor, limit) for mark in (";", ",", " "))
            if split_at <= cursor + CHUNK_SIZE // 2:
                split_at = limit
            else:
                split_at += 1
            bounded.append((cursor, split_at))
            cursor = split_at
        if cursor < sentence_end:
            bounded.append((cursor, sentence_end))
    return bounded


def _pack_chunk_ranges(
    text: str,
    unit_ranges: List[Tuple[int, int]],
) -> List[Tuple[int, int]]:
    """Pack adjacent legal units without crossing the target size or input span."""

    packed = []
    chunk_start = chunk_end = None
    for unit_start, unit_end in unit_ranges:
        if chunk_start is None:
            chunk_start, chunk_end = unit_start, unit_end
        elif unit_end - chunk_start > CHUNK_SIZE:
            packed.append((chunk_start, chunk_end))
            chunk_start, chunk_end = unit_start, unit_end
        else:
            chunk_end = unit_end
    if chunk_start is not None and chunk_end is not None:
        packed.append((chunk_start, chunk_end))
    return packed


def _semantic_chunk_ranges(
    text: str,
    unit_ranges: List[Tuple[int, int]],
) -> Optional[Tuple[List[Tuple[int, int]], bool]]:
    """Split at unusually large semantic shifts without crossing legal units.

    Units are paragraph-sized where possible and sentence-sized only when a
    paragraph exceeds the normal chunk limit. The already configured passage
    embedder is used for boundary detection; callers retain the structural
    packer as a fallback if the model is unavailable or the document is too
    large for a bounded semantic pass.
    """

    if (
        not getattr(settings, "SEMANTIC_CHUNKING_ENABLED", True)
        or len(unit_ranges) < 2
        or unit_ranges[-1][1] - unit_ranges[0][0] <= CHUNK_SIZE
    ):
        return None

    try:
        maximum_units = max(2, min(512, int(getattr(settings, "SEMANTIC_CHUNK_MAX_UNITS", 160))))
        if len(unit_ranges) > maximum_units:
            return None
        passages = [text[start:end].strip() for start, end in unit_ranges]
        if any(not passage for passage in passages):
            return None

        vectors = embed_texts(passages, "passage")
        if len(vectors) != len(passages):
            return None

        distances = []
        for left, right in zip(vectors, vectors[1:]):
            if len(left) != len(right) or not left:
                return None
            dot = sum(float(a) * float(b) for a, b in zip(left, right))
            left_norm = math.sqrt(sum(float(value) ** 2 for value in left))
            right_norm = math.sqrt(sum(float(value) ** 2 for value in right))
            if left_norm <= 0 or right_norm <= 0:
                return None
            similarity = max(-1.0, min(1.0, dot / (left_norm * right_norm)))
            distances.append(1.0 - similarity)

        percentile = max(
            50.0,
            min(100.0, float(getattr(settings, "SEMANTIC_CHUNK_DISTANCE_PERCENTILE", 90.0))),
        )
        ordered_distances = sorted(distances)
        percentile_index = min(
            len(ordered_distances) - 1,
            max(0, math.ceil(percentile / 100.0 * len(ordered_distances)) - 1),
        )
        threshold = max(
            0.0,
            float(getattr(settings, "SEMANTIC_CHUNK_MIN_DISTANCE", 0.16)),
            ordered_distances[percentile_index],
        )
        minimum_chars = max(100, min(1000, int(getattr(settings, "SEMANTIC_CHUNK_MIN_CHARS", 320))))

        ranges: List[Tuple[int, int]] = []
        first_unit = 0
        semantic_boundary_used = False
        for next_unit in range(1, len(unit_ranges)):
            start = unit_ranges[first_unit][0]
            previous_end = unit_ranges[next_unit - 1][1]
            next_end = unit_ranges[next_unit][1]
            forced_boundary = next_end - start > CHUNK_SIZE
            remaining_chars = unit_ranges[-1][1] - unit_ranges[next_unit][0]
            semantic_boundary = (
                distances[next_unit - 1] >= threshold
                and previous_end - start >= minimum_chars
                and remaining_chars >= minimum_chars
            )
            if not forced_boundary and not semantic_boundary:
                continue
            ranges.append((start, previous_end))
            semantic_boundary_used = semantic_boundary_used or semantic_boundary
            first_unit = next_unit

        ranges.append((unit_ranges[first_unit][0], unit_ranges[-1][1]))
        return ranges, semantic_boundary_used
    except Exception as exc:
        # Extraction must remain usable when embeddings are unavailable. The
        # fallback is the existing deterministic structural chunker.
        logger.debug("Semantic chunk boundary pass unavailable: %s", type(exc).__name__)
        return None


def clean_article_title(title: str) -> str:
    title = re.sub(r"\s+", " ", title or "").strip()
    title = re.sub(r"\.\d+\s*$", "", title)
    return title.strip(" .:-")


def extract_article_refs(text: str) -> List[Tuple[int, Dict]]:
    refs = []
    for match in ARTICLE_RE.finditer(text):
        article_number = match.group("number").strip()
        article_title = clean_article_title(match.group("title"))
        citation = f"ART. {article_number}"
        if article_title:
            citation = f"{citation}. {article_title}"
        refs.append((
            match.start(),
            {
                "article_number": article_number,
                "article_title": article_title,
                "citation": citation,
            },
        ))
    return refs


def metadata_for_chunk(
    source: str,
    page: int,
    article_refs: List[Tuple[int, Dict]],
    start: int,
    end: int,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> Dict:
    metadata = {
        "source": source,
        "page": page,
        "category": normalize_kb_category(category),
    }
    article_meta = {}

    for position, ref in article_refs:
        if position <= start:
            article_meta = ref
            continue
        if position >= end:
            break

    if article_meta:
        metadata.update(article_meta)
    return metadata


def article_parent_chunks(
    text: str,
    source: str,
    page: int = 0,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> List[Tuple[str, Dict]]:
    text = clean_chunk_text(text)
    article_refs = extract_article_refs(text)
    if not article_refs:
        return []

    chunks: List[Tuple[str, Dict]] = []
    for index, (start, article_meta) in enumerate(article_refs):
        end = article_refs[index + 1][0] if index + 1 < len(article_refs) else len(text)
        if end - start < 80:
            continue
        units = [
            unit
            for paragraph_start, paragraph_end in _paragraph_ranges(text, start, end)
            for unit in _bounded_unit_ranges(text, paragraph_start, paragraph_end)
        ]
        ranges = _pack_chunk_ranges(text, units)
        for part_index, (part_start, part_end) in enumerate(ranges):
            article_text = text[part_start:part_end].strip()
            if not article_text:
                continue
            metadata = {
                "source": source,
                "page": page,
                "category": normalize_kb_category(category),
                "chunk_role": "article_parent",
                "article_chunk_index": part_index,
                "article_chunk_count": len(ranges),
                **article_meta,
            }
            chunks.append((article_text, metadata))
    return chunks


def document_article_parent_chunks(
    pages: List[Dict],
    source: str,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> List[Tuple[str, Dict]]:
    """Build complete article chunks across PDF page boundaries."""
    combined_parts: List[str] = []
    page_offsets: List[Tuple[int, int]] = []
    cursor = 0
    for item in sorted(pages, key=lambda value: int(value.get("page") or 0)):
        page_text = clean_chunk_text(item.get("text", ""))
        if not page_text:
            continue
        if combined_parts:
            cursor += 2
        page_offsets.append((cursor, int(item.get("page") or 0)))
        combined_parts.append(page_text)
        cursor += len(page_text)

    combined = "\n\n".join(combined_parts)
    article_refs = extract_article_refs(combined)
    chunks: List[Tuple[str, Dict]] = []
    for index, (start, article_meta) in enumerate(article_refs):
        end = article_refs[index + 1][0] if index + 1 < len(article_refs) else len(combined)
        structural_boundary = STRUCTURAL_BOUNDARY_RE.search(combined, start + 40, end)
        if structural_boundary:
            end = structural_boundary.start()
        if end - start < 80:
            continue
        units = [
            unit
            for paragraph_start, paragraph_end in _paragraph_ranges(combined, start, end)
            for unit in _bounded_unit_ranges(combined, paragraph_start, paragraph_end)
        ]
        ranges = _pack_chunk_ranges(combined, units)
        for part_index, (part_start, part_end) in enumerate(ranges):
            article_text = combined[part_start:part_end].strip()
            if not article_text:
                continue
            covered_pages = [
                page for offset, page in page_offsets if part_start <= offset < part_end
            ]
            start_page = max(
                (page for offset, page in page_offsets if offset <= part_start),
                default=0,
            )
            end_page = max(covered_pages, default=start_page)
            chunks.append((
                article_text,
                {
                    "source": source,
                    "page": start_page,
                    "page_end": end_page,
                    "category": normalize_kb_category(category),
                    "chunk_role": "document_article_parent",
                    "article_chunk_index": part_index,
                    "article_chunk_count": len(ranges),
                    **article_meta,
                },
            ))
    return chunks


def document_section_parent_chunks(pages, source, category=DEFAULT_KB_CATEGORY):
    """Index section parts across pages, including hierarchy and stable order."""
    from .section_structure import section_ranges
    combined_parts, offsets, cursor = [], [], 0
    for item in sorted(pages, key=lambda value: int(value.get('page') or 0)):
        text = clean_chunk_text(item.get('text') or '')
        if not text:
            continue
        if combined_parts:
            cursor += 2
        offsets.append((cursor, int(item.get('page') or 0)))
        combined_parts.append(text)
        cursor += len(text)
    combined = '\n\n'.join(combined_parts)
    chunks = []
    for start, end, locator in section_ranges(combined):
        units = [unit for a, b in _paragraph_ranges(combined, start, end)
                 for unit in _bounded_unit_ranges(combined, a, b)]
        ranges = _pack_chunk_ranges(combined, units)
        for index, (a, b) in enumerate(ranges):
            page = max((p for offset, p in offsets if offset <= a), default=0)
            page_end = max((p for offset, p in offsets if a <= offset < b), default=page)
            chunks.append((combined[a:b].strip(), {
                'source': source, 'page': page, 'page_end': page_end,
                'category': normalize_kb_category(category),
                'chunk_role': 'document_section_parent', 'section_part_index': index,
                'section_part_count': len(ranges), **locator,
            }))
    return chunks


def chunk_text(
    text: str,
    source: str,
    page: int = 0,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
    include_article_parents: bool = True,
    leading_article_metadata: Optional[Dict] = None,
) -> List[Tuple[str, Dict]]:
    chunks = []
    text = clean_chunk_text(text)
    if not text:
        return chunks
    article_refs = extract_article_refs(text)
    parent_chunks = (
        article_parent_chunks(text, source=source, page=page, category=category)
        if include_article_parents
        else []
    )

    if len(text) <= CHUNK_SIZE and not (
        leading_article_metadata and article_refs and article_refs[0][0] > 0
    ):
        base_metadata = metadata_for_chunk(source, page, article_refs, 0, len(text), category)
        if leading_article_metadata and not base_metadata.get("article_number"):
            base_metadata.update(leading_article_metadata)
        base_chunks = [(
            text,
            base_metadata,
        )]
        return parent_chunks + [
            item for item in base_chunks
            if not parent_chunks or item[1].get("citation") != parent_chunks[0][1].get("citation")
        ]

    # Keep separate articles from being packed together, even if each is short.
    article_starts = [position for position, _ in article_refs]
    group_starts = sorted(set([0, *article_starts]))
    group_ranges = [
        (start, group_starts[index + 1] if index + 1 < len(group_starts) else len(text))
        for index, start in enumerate(group_starts)
    ]

    for group_start, group_end in group_ranges:
        units = []
        for paragraph_start, paragraph_end in _paragraph_ranges(text, group_start, group_end):
            units.extend(_bounded_unit_ranges(text, paragraph_start, paragraph_end))
        if not units:
            continue
        semantic_result = _semantic_chunk_ranges(text, units)
        if semantic_result is None:
            chunk_ranges = _pack_chunk_ranges(text, units)
            used_semantic_boundary = False
        else:
            chunk_ranges, used_semantic_boundary = semantic_result

        for chunk_start, chunk_end in chunk_ranges:
            chunk = text[chunk_start:chunk_end].strip()
            if chunk:
                chunk_metadata = metadata_for_chunk(
                    source, page, article_refs, chunk_start, chunk_end, category
                )
                if leading_article_metadata and not chunk_metadata.get("article_number"):
                    chunk_metadata.update(leading_article_metadata)
                chunk_metadata["chunking_method"] = (
                    "semantic" if used_semantic_boundary else "structural"
                )
                chunks.append((chunk, chunk_metadata))

    return parent_chunks + chunks


def mark_extraction_method(
    chunks: List[Tuple[str, Dict]],
    method: str,
) -> List[Tuple[str, Dict]]:
    return [
        (text, {**metadata, "extraction_method": method})
        for text, metadata in chunks
    ]


def deduplicate_ocr_chunks(
    chunks: List[Tuple[str, Dict]],
) -> List[Tuple[str, Dict]]:
    """Remove repeated OCR chunks while leaving normal text chunks untouched.

    Scanned PDFs often repeat page images, headers, or OCR text blocks. We use a
    conservative normalized full-text fingerprint, so overlapping but distinct
    legal passages remain available to retrieval.
    """

    unique_chunks: List[Tuple[str, Dict]] = []
    seen_ocr_text: set[str] = set()
    for text, metadata in chunks:
        if metadata.get("extraction_method") != "ocr":
            unique_chunks.append((text, metadata))
            continue

        fingerprint = re.sub(r"\W+", " ", text.casefold()).strip()
        if fingerprint and fingerprint in seen_ocr_text:
            continue
        if fingerprint:
            seen_ocr_text.add(fingerprint)
        unique_chunks.append((text, metadata))
    return unique_chunks


def tesseract_cmd() -> Optional[str]:
    configured = (getattr(settings, "TESSERACT_CMD", "") or "").strip().strip('"')

    if configured:
        configured_path = Path(configured).expanduser()
        if configured_path.is_file():
            return str(configured_path)
        configured_from_path = shutil.which(configured)
        if configured_from_path:
            return configured_from_path

    path_match = shutil.which("tesseract")
    if path_match:
        return path_match

    common_windows_paths = [
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
        Path.home() / "AppData" / "Local" / "Programs" / "Tesseract-OCR" / "tesseract.exe",
    ]

    for candidate in common_windows_paths:
        if candidate.is_file():
            return str(candidate)

    return None


def ensure_ocr_available(require_pdf_renderer: bool = False) -> None:
    missing = []
    if Image is None:
        missing.append("Pillow")
    if tesseract_cmd() is None:
        missing.append("Tesseract OCR engine")
    if (
        require_pdf_renderer
        and convert_from_path is None
        and shutil.which("pdftoppm") is None
        and pdfium is None
    ):
        missing.append("pdf2image, Poppler, or pypdfium2")
    if missing:
        raise RuntimeError(
            "OCR support is not installed. Install "
            + ", ".join(missing)
            + " in the backend environment or server."
        )


def ordered_ocr_text(tsv: str) -> str:
    """Keep verbatim OCR words in geometric reading order, including tables.

    Tesseract's block traversal may emit table values before the introduction.
    Coordinates restore top-to-bottom order; large horizontal gaps retain a
    cell separator instead of silently joining unrelated numeric columns.
    """
    words = []
    for row in csv.DictReader(StringIO(tsv), delimiter='\t', quoting=csv.QUOTE_NONE):
        if str(row.get('level')) != '5' or not (row.get('text') or '').strip():
            continue
        try:
            words.append({'text': row['text'].strip(), 'x': int(row['left']),
                          'y': int(row['top']) + int(row['height']) / 2,
                          'width': int(row['width']), 'height': int(row['height']),
                          'paragraph': (row.get('block_num'), row.get('par_num'))})
        except (ValueError, TypeError, KeyError):
            continue
    rows = []
    for word in sorted(words, key=lambda w: (w['y'], w['x'])):
        if rows and abs(word['y'] - rows[-1]['y']) <= max(3, min(
                word['height'], rows[-1]['height']) * .6):
            rows[-1]['words'].append(word)
        else:
            rows.append({'y': word['y'], 'height': word['height'], 'words': [word]})
    output, previous = [], None
    for row in rows:
        line, last = '', None
        ordered = sorted(row['words'], key=lambda w: w['x'])
        for word in ordered:
            gap = word['x'] - (last['x'] + last['width']) if last else 0
            separator = ' | ' if last and gap > max(20, row['height'] * 1.8) else ' '
            line += (separator if last else '') + word['text']
            last = word
        if previous and (row['y'] - previous['y'] > max(row['height'], previous['height']) * 1.8
                         or ordered[0]['paragraph'] != previous['words'][0]['paragraph']):
            output.append('')
        output.append(line)
        previous = row
    return '\n'.join(output)


def recover_ocr_table_cells(image, tsv, cmd):
    """Retry damaged currency rows as a bounded table band, without guessing.

    Replace words geometrically, preserving the complete page and cell labels.
    A retry is accepted only if all suspect cells remain and become numeric.
    """
    from .extraction_quality import extraction_issues
    rows = list(csv.DictReader(StringIO(tsv), delimiter='\t', quoting=csv.QUOTE_NONE))
    bad = [r for r in rows if r.get('level') == '5' and
           'suspicious_currency_ocr' in extraction_issues(r.get('text'))]
    if not bad:
        return tsv
    top = max(0, min(int(r['top']) for r in bad) - max(80, int(image.height * .06)))
    bottom = min(image.height, max(int(r['top']) + int(r['height']) for r in bad) + 40)
    cropped = image.crop((0, top, image.width, bottom))
    try:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'table-band.png'
            cropped.save(path)
            retry = subprocess.run([cmd, str(path), 'stdout', '--psm', '11', 'tsv'],
                capture_output=True, text=True, encoding='utf-8', errors='replace',
                timeout=60, env={**os.environ, 'OMP_THREAD_LIMIT': '1'})
        if retry.returncode != 0:
            return tsv
        replacement = [r for r in csv.DictReader(StringIO(retry.stdout), delimiter='\t',
                                                quoting=csv.QUOTE_NONE) if r.get('level') == '5']
        old_band = [r for r in rows if r.get('level') == '5' and
                    top <= int(r['top']) + int(r['height']) / 2 <= bottom]
        old_amounts = [r for r in old_band if re.search(r'\bphp|₱', r.get('text', ''), re.I)]
        new_amounts = [r for r in replacement if re.search(r'(?:\bph?p|₱)\s*\d[\d,]*(?:\.\d+)?', r.get('text', ''), re.I)]
        if (len(new_amounts) < len(old_amounts) or
                any(extraction_issues(r.get('text')) for r in replacement)):
            return tsv
        kept = [r for r in rows if r.get('level') == '5' and r not in old_band]
        for r in replacement:
            r['top'] = str(int(r['top']) + top)
            r['block_num'] = 'recovered_table'
        stream = StringIO()
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter='\t')
        writer.writeheader()
        writer.writerows(kept + replacement)
        return stream.getvalue()
    except (ValueError, KeyError, subprocess.TimeoutExpired):
        return tsv
    finally:
        cropped.close()


def ocr_image_to_text(image, label: str = "image") -> str:
    ensure_ocr_available()
    cmd = tesseract_cmd()
    prepared = None
    try:
        prepared = image.convert("L")
        with tempfile.TemporaryDirectory() as tmpdir:
            image_path = Path(tmpdir) / "ocr-input.png"
            prepared.save(image_path)
            # Recover sideways scanned tables before reading words. Orientation
            # failure is nonfatal; it must never erase an otherwise readable page.
            try:
                orientation = subprocess.run([cmd, str(image_path), 'stdout', '--psm', '0'],
                    capture_output=True, text=True, encoding='utf-8', errors='replace',
                    timeout=20, env={**os.environ, 'OMP_THREAD_LIMIT': '1'})
                angle = re.search(r'Rotate:\s*(90|180|270)\b', orientation.stdout)
                confidence = re.search(r'Orientation confidence:\s*([\d.]+)', orientation.stdout)
                if angle and confidence and float(confidence.group(1)) >= 5:
                    rotated = prepared.rotate(-int(angle.group(1)), expand=True)
                    prepared.close()
                    prepared = rotated
                    prepared.save(image_path)
            except (subprocess.TimeoutExpired, ValueError):
                logger.debug('OCR orientation check unavailable on %s', label)
            result = subprocess.run(
                [cmd, str(image_path), "stdout", "--psm", "3", "tsv"],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
                timeout=120, env={**os.environ, "OMP_THREAD_LIMIT": "1"},
            )
            if result.returncode != 0:
                raise RuntimeError(result.stderr.strip() or "Tesseract OCR failed.")
            text = ordered_ocr_text(recover_ocr_table_cells(prepared, result.stdout, cmd))
    except Exception as exc:
        raise RuntimeError(f"OCR failed on {label}: {exc}. Retry to reuse completed PDF pages.") from exc
    finally:
        if prepared is not None:
            prepared.close()
    text = _normalize_document_text(text)
    logger.info("OCR extracted %s characters from %s", len(text), label)
    return text


def render_pdf_page_with_pdftoppm(filepath: str, page_num: int):
    if shutil.which("pdftoppm") is None:
        raise RuntimeError("Poppler pdftoppm is not available on the server PATH.")
    with tempfile.TemporaryDirectory() as tmpdir:
        output_prefix = Path(tmpdir) / "page"
        result = subprocess.run(
            [
                "pdftoppm",
                "-f",
                str(page_num),
                "-l",
                str(page_num),
                "-r",
                str(OCR_DPI),
                "-png",
                filepath,
                str(output_prefix),
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or "pdftoppm failed.")
        rendered = sorted(Path(tmpdir).glob("page-*.png"))
        if not rendered:
            raise RuntimeError("pdftoppm did not render a page image.")
        with Image.open(rendered[0]) as image:
            return image.copy()


def render_pdf_page_with_pdfium(filepath: str, page_num: int):
    global pdfium
    if pdfium is None:
        try:
            import pypdfium2 as loaded_pdfium
            pdfium = loaded_pdfium
        except ImportError as exc:
            raise RuntimeError("pypdfium2 is not installed.") from exc
    pdf = None
    page = None
    bitmap = None
    image = None
    try:
        pdf = pdfium.PdfDocument(filepath)
        page = pdf[page_num - 1]
        bitmap = page.render(scale=OCR_DPI / 72)
        image = bitmap.to_pil().convert("RGB")
        return image.copy()
    except Exception as exc:
        raise RuntimeError("pypdfium2 failed to render the scanned PDF page.") from exc
    finally:
        if image is not None:
            image.close()
        for resource in (bitmap, page, pdf):
            close = getattr(resource, "close", None)
            if callable(close):
                try:
                    close()
                except Exception:
                    pass


def render_pdf_page_for_ocr(filepath: str, page_num: int, grayscale: bool = True):
    errors = []

    if convert_from_path is not None:
        try:
            images = convert_from_path(
                filepath,
                dpi=OCR_DPI,
                first_page=page_num,
                last_page=page_num,
                thread_count=1,
                grayscale=grayscale,
                timeout=120,
            )
            if images:
                return images[0]
            errors.append("pdf2image did not return an image")
        except Exception as exc:
            errors.append(f"pdf2image failed: {exc}")

    if shutil.which("pdftoppm") is not None:
        try:
            return render_pdf_page_with_pdftoppm(filepath, page_num)
        except Exception as exc:
            errors.append(f"Poppler pdftoppm failed: {exc}")

    try:
        return render_pdf_page_with_pdfium(filepath, page_num)
    except Exception as exc:
        errors.append(f"pypdfium2 failed: {exc}")

    detail = "; ".join(errors) if errors else "no PDF renderer is available"
    raise RuntimeError(
        "PDF-to-image conversion failed. "
        f"{detail}. Install pypdfium2 or Poppler so scanned PDF pages can be converted before OCR."
    )


def ocr_pdf_pages(filepath: str, page_numbers: Optional[List[int]] = None) -> List[Dict]:
    ensure_ocr_available(require_pdf_renderer=True)
    reader = pdf_reader_from_file(filepath)
    total_pages = len(reader.pages)
    pages_to_scan = page_numbers if page_numbers is not None else list(range(1, total_pages + 1))
    ocr_pages = []
    checkpoint = OCRCheckpoint(filepath, f"ocr-v4-oriented-geometric-tsv-psm3-dpi{OCR_DPI}")
    previous_checkpoint = OCRCheckpoint(filepath, f"ocr-v3-geometric-tsv-psm3-dpi{OCR_DPI}")

    for page_num in pages_to_scan:
        if page_num < 1 or page_num > total_pages:
            continue
        cached = checkpoint.read(page_num)
        if cached is None:
            from .extraction_quality import extraction_issues
            previous_text = previous_checkpoint.read(page_num)
            if previous_text and not extraction_issues(previous_text):
                cached = previous_text
        if cached is not None:
            if cached:
                ocr_pages.append({"page": page_num, "text": cached})
            logger.info("Reused OCR page %s/%s", page_num, total_pages)
            continue
        image = None
        # Bound image memory and OCR contention when several scans are uploaded.
        with _OCR_PAGE_LOCK:
            try:
                image = render_pdf_page_for_ocr(filepath, page_num)
                text = ocr_image_to_text(image, f"page {page_num}/{total_pages}")
                try:
                    checkpoint.write(page_num, text)
                except OSError:
                    logger.warning("Could not checkpoint OCR page %s", page_num)
                if text:
                    ocr_pages.append({"page": page_num, "text": text})
            finally:
                close = getattr(image, "close", None)
                if callable(close):
                    close()
    return ocr_pages


def extract_from_image(
    filepath: str,
    source_name: str = None,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> List[Tuple[str, Dict]]:
    ensure_ocr_available()
    filename = source_name or Path(filepath).name
    try:
        with Image.open(filepath) as image:
            text = ocr_image_to_text(image, filename)
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError("Unable to open this image for OCR.") from exc

    chunks = chunk_text(text, source=filename, page=1, category=category)
    chunks = mark_extraction_method(chunks, "ocr")
    logger.info("Extracted %s OCR chunks from image %s", len(chunks), filename)
    return chunks


def chunks_from_ocr_pages(
    pages: List[Dict],
    source: str,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> List[Tuple[str, Dict]]:
    all_chunks = []
    for item in pages:
        page_num = item.get("page", 0)
        text = item.get("text", "")
        page_chunks = chunk_text(text, source=source, page=page_num, category=category)
        all_chunks.extend(mark_extraction_method(page_chunks, "ocr"))
    return all_chunks


def _pdf_matrix_product(current, local):
    """Compose PDF affine transforms, including nested image/form placement."""
    a, b, c, d, e, f = current
    g, h, i, j, k, l = local
    return (
        a * g + c * h, b * g + d * h,
        a * i + c * j, b * i + d * j,
        a * k + c * l + e, b * k + d * l + f,
    )


def _resolve_pdf_object(value):
    """Dereference nested pypdf indirect objects without looping forever."""
    for _ in range(12):
        get_object = getattr(value, "get_object", None)
        if not callable(get_object):
            break
        resolved = get_object()
        if resolved is value:
            break
        value = resolved
    return value


def _pdf_has_substantial_image_resource(page) -> bool:
    """Detect page-sized image resources when malformed content streams omit Do."""
    resources = _resolve_pdf_object(page.get("/Resources", {})) or {}
    xobjects = _resolve_pdf_object(resources.get("/XObject", {})) or {}
    for reference in xobjects.values():
        image = _resolve_pdf_object(reference)
        if image.get("/Subtype") != "/Image":
            continue
        width = int(_resolve_pdf_object(image.get("/Width", 0)))
        height = int(_resolve_pdf_object(image.get("/Height", 0)))
        if width * height >= _PDF_IMAGE_MIN_PIXELS:
            return True
    return False


def _pdf_image_regions(page) -> List[Tuple[float, float, float, float]]:
    """Find substantial placed raster images, including images inside forms.

    Text extraction alone cannot distinguish a decorative logo from an image
    containing a table. Placement and pixel size provide a conservative signal
    for OCR, rather than treating the presence of any native text as complete.
    """
    box = page.cropbox
    page_box = tuple(float(value) for value in (box.left, box.bottom, box.right, box.top))
    page_area = max(1.0, (page_box[2] - page_box[0]) * (page_box[3] - page_box[1]))
    regions = []

    def record_image(width, height, image_mask, matrix):
        width = int(_resolve_pdf_object(width))
        height = int(_resolve_pdf_object(height))
        image_mask = _resolve_pdf_object(image_mask)
        if width * height < _PDF_IMAGE_MIN_PIXELS or image_mask:
            return
        a, b, c, d, e, f = matrix
        points = [(e, f), (a + e, b + f), (c + e, d + f), (a + c + e, b + d + f)]
        left = max(page_box[0], min(point[0] for point in points))
        bottom = max(page_box[1], min(point[1] for point in points))
        right = min(page_box[2], max(point[0] for point in points))
        top = min(page_box[3], max(point[1] for point in points))
        if right <= left or top <= bottom:
            return
        area = (right - left) * (top - bottom)
        if area / page_area >= _PDF_IMAGE_MIN_PAGE_AREA:
            regions.append((left, bottom, right, top))
            if len(regions) > _PDF_IMAGE_MAX_REGIONS:
                raise RuntimeError("PDF image inspection exceeded the placed-image limit.")

    def visit(stream, resources, transform, ancestors, depth):
        if stream is None:
            return
        if depth > _PDF_IMAGE_MAX_FORM_DEPTH:
            raise RuntimeError("PDF image inspection exceeded the nested-form limit.")
        content = stream if isinstance(stream, ContentStream) else ContentStream(stream, page.pdf)
        matrix = transform
        stack = []
        resources = _resolve_pdf_object(resources) or {}
        xobjects = _resolve_pdf_object(resources.get("/XObject", {})) or {}
        for operands, operator in content.operations:
            if operator == b"q":
                stack.append(matrix)
            elif operator == b"Q":
                matrix = stack.pop() if stack else transform
            elif operator == b"cm" and len(operands) == 6:
                matrix = _pdf_matrix_product(matrix, tuple(float(value) for value in operands))
            elif operator == b"INLINE IMAGE":
                image_settings = operands.get("settings", {})
                record_image(
                    image_settings.get("/W", image_settings.get("/Width", 0)),
                    image_settings.get("/H", image_settings.get("/Height", 0)),
                    image_settings.get("/IM", image_settings.get("/ImageMask", False)), matrix,
                )
            elif operator == b"Do" and operands:
                item = xobjects.get(operands[0])
                if item is None:
                    continue
                item = item.get_object()
                if item.get("/Subtype") == "/Form":
                    identity = id(item)
                    if identity in ancestors:
                        raise RuntimeError("PDF contains a circular image/form reference.")
                    form_matrix = tuple(float(value) for value in item.get("/Matrix", (1, 0, 0, 1, 0, 0)))
                    visit(
                        item, item.get("/Resources", resources),
                        _pdf_matrix_product(matrix, form_matrix),
                        ancestors | {identity}, depth + 1,
                    )
                elif item.get("/Subtype") == "/Image":
                    record_image(item.get("/Width", 0), item.get("/Height", 0), item.get("/ImageMask"), matrix)

    visit(page.get_contents(), page.get("/Resources", {}), (1, 0, 0, 1, 0, 0), set(), 0)
    return regions


def _pdf_native_page_text(page) -> Tuple[str, bool]:
    """Return native text and whether visible image content needs an OCR pass."""
    positioned_text = []

    def collect_text(text, current_matrix, text_matrix, _font, _font_size):
        if text and text.strip():
            a, b, c, d, e, f = current_matrix
            x, y = float(text_matrix[4]), float(text_matrix[5])
            positioned_text.append((a * x + c * y + e, b * x + d * y + f, text))

    text = _normalize_document_text(page.extract_text(visitor_text=collect_text) or "")
    from .extraction_quality import extraction_issues
    if extraction_issues(text):
        # A large text layer can itself be an old, damaged OCR pass.
        # Character count alone cannot certify extraction completeness.
        return text, True
    try:
        regions = _pdf_image_regions(page)
    except Exception as exc:
        logger.warning("PDF image inspection requires an OCR fallback: %s", exc)
        return text, True
    if not text:
        contents = page.get_contents()
        # Genuine blank pages have no evidence to recover. Outlined lettering
        # and drawing-based content can require OCR even without image objects.
        painted = contents is not None and any(
            operator in {b"Tj", b"TJ", b"f", b"F", b"f*", b"S", b"s", b"B", b"B*", b"b", b"b*", b"INLINE IMAGE"}
            for _operands, operator in contents.operations
        )
        return text, bool(regions or painted or _pdf_has_substantial_image_resource(page))
    for left, bottom, right, top in regions:
        visible_chars = sum(
            len(re.sub(r"\s+", "", part))
            for x, y, part in positioned_text
            if left - 5 <= x <= right + 5 and bottom - 12 <= y <= top + 12
        )
        # A native OCR layer covering the image avoids a redundant OCR pass.
        # A caption or a few labels over a large raster area is not sufficient.
        required_chars = min(1200, max(40, int((right - left) * (top - bottom) / 450)))
        if visible_chars < required_chars:
            return text, True
    return text, False


def _merge_pdf_native_and_ocr(native_text: str, ocr_text: str) -> str:
    """Retain exact native text and complementary OCR blocks, with table context.

    Deduplicate complete blocks only. Removing individual repeated table
    headings can detach newly recovered values from their labels and units.
    """
    native_text = _normalize_document_text(native_text)
    ocr_text = _normalize_document_text(ocr_text)
    if not native_text:
        return ocr_text
    native_key = re.sub(r"\s+", " ", native_text).casefold()
    additional = []
    seen = set()
    for block in re.split(r"\n\s*\n", ocr_text):
        key = re.sub(r"\s+", " ", block).strip().casefold()
        if key and key not in native_key and key not in seen:
            additional.append(block)
            seen.add(key)
    return "\n\n".join([native_text, *additional]).strip()


def _is_nearly_blank_page(image) -> bool:
    """Treat a nearly white page with only scan specks as non-content."""
    if Image is None or image is None:
        return False
    grayscale = None
    dark_mask = None
    try:
        grayscale = image.convert("L")
        grayscale.thumbnail((256, 256))
        dark_mask = grayscale.point(lambda value: 255 if value < 220 else 0)
        dark_pixels = dark_mask.histogram()[255]
        return dark_pixels / max(1, grayscale.width * grayscale.height) <= 0.001
    except Exception as exc:
        logger.debug("Could not classify an empty OCR page as nearly blank: %s", exc)
        return False
    finally:
        if dark_mask is not None:
            dark_mask.close()
        if grayscale is not None:
            grayscale.close()


def _is_blank_ruled_page(image) -> bool:
    """Recognize an empty form page made only of evenly spaced writing lines."""
    if Image is None or image is None:
        return False
    grayscale = None
    dark_mask = None
    try:
        grayscale = image.convert("L")
        grayscale.thumbnail((512, 512))
        width, height = grayscale.size
        dark_mask = grayscale.point(lambda value: 255 if value < 220 else 0)
        pixels = dark_mask.load()
        row_counts = [sum(1 for x in range(width) if pixels[x, y]) for y in range(height)]
        rule_rows = [y for y, count in enumerate(row_counts) if count >= width * 0.68]
        if len(rule_rows) < 12:
            return False

        groups = []
        for row in rule_rows:
            if not groups or row - groups[-1][-1] > 1:
                groups.append([row])
            else:
                groups[-1].append(row)
        centers = [sum(group) / len(group) for group in groups]
        gaps = [right - left for left, right in zip(centers, centers[1:])]
        if len(gaps) < 10:
            return False
        median_gap = sorted(gaps)[len(gaps) // 2]
        if not 3 <= median_gap <= 24:
            return False
        if sum(row_counts[row] for row in rule_rows) / max(1, sum(row_counts)) < 0.85:
            return False
        return all(abs(gap - median_gap) <= max(1.5, median_gap * 0.25) for gap in gaps)
    except Exception as exc:
        logger.debug("Could not classify an empty OCR page as a blank ruled form: %s", exc)
        return False
    finally:
        if dark_mask is not None:
            dark_mask.close()
        if grayscale is not None:
            grayscale.close()


def _is_centered_color_graphic_page(image) -> bool:
    """Recognize sparse centered artwork (such as an agency logo) without text.

    This is only used after OCR returns no text. It requires both a compact,
    saturated-color region and a compact dark-content region so a failed OCR
    pass over a typical text page still fails closed.
    """
    if Image is None or ImageChops is None or image is None:
        return False

    rgb = None
    channels = []
    masks = []
    intermediates = []
    try:
        rgb = image.convert("RGB")
        rgb.thumbnail((256, 256))
        if not rgb.width or not rgb.height:
            return False

        channels = list(rgb.split())
        maximum = ImageChops.lighter(ImageChops.lighter(channels[0], channels[1]), channels[2])
        minimum = ImageChops.darker(ImageChops.darker(channels[0], channels[1]), channels[2])
        saturation = ImageChops.difference(maximum, minimum)
        color_mask = saturation.point(lambda value: 255 if value >= 45 else 0)
        grayscale = rgb.convert("L")
        dark_mask = grayscale.point(lambda value: 255 if value < 100 else 0)
        masks.extend((maximum, minimum, saturation, color_mask, grayscale, dark_mask))

        width, height = rgb.size
        total_pixels = width * height
        color_ratio = color_mask.histogram()[255] / total_pixels
        dark_ratio = dark_mask.histogram()[255] / total_pixels
        color_box = color_mask.getbbox()
        dark_box = dark_mask.getbbox()

        # Some back covers use full-page blocks of flat color. OCR correctly
        # finds no text, but the artwork is not a compact centered logo.
        full_page_flat_art = False
        if color_ratio >= 0.80:
            small = rgb.resize((128, 128), Image.NEAREST)
            small_gray = small.convert("L")
            left = small_gray.crop((0, 0, 127, 128))
            right = small_gray.crop((1, 0, 128, 128))
            top = small_gray.crop((0, 0, 128, 127))
            bottom = small_gray.crop((0, 1, 128, 128))
            horizontal_edges = ImageChops.difference(left, right).point(
                lambda value: 255 if value > 20 else 0
            )
            vertical_edges = ImageChops.difference(top, bottom).point(
                lambda value: 255 if value > 20 else 0
            )
            intermediates.extend((small, small_gray, left, right, top, bottom))
            masks.extend((horizontal_edges, vertical_edges))
            horizontal_ratio = horizontal_edges.histogram()[255] / (127 * 128)
            vertical_ratio = vertical_edges.histogram()[255] / (128 * 127)
            full_page_flat_art = horizontal_ratio <= 0.025 and vertical_ratio <= 0.025

        if full_page_flat_art:
            return True
        if not color_box or not dark_box:
            return False

        def compact_and_centered(box, max_width, max_height):
            left, top, right, bottom = box
            box_width = (right - left) / width
            box_height = (bottom - top) / height
            center_x = ((left + right) / 2) / width
            center_y = ((top + bottom) / 2) / height
            return (
                box_width <= max_width
                and box_height <= max_height
                and abs(center_x - 0.5) <= 0.20
                and abs(center_y - 0.5) <= 0.20
            )

        return (
            0.012 <= color_ratio <= 0.12
            and 0.005 <= dark_ratio <= 0.10
            and compact_and_centered(color_box, 0.55, 0.45)
            and compact_and_centered(dark_box, 0.60, 0.55)
        )
    except Exception as exc:
        logger.debug("Could not classify an empty OCR page as decorative: %s", exc)
        return False
    finally:
        for item in masks + channels + intermediates:
            close = getattr(item, "close", None)
            if callable(close):
                close()
        if rgb is not None:
            rgb.close()


def _extract_pdf_page_records(filepath: str) -> Tuple[List[Dict], int]:
    """Share complete extraction behavior between indexing and text preview."""
    reader = pdf_reader_from_file(filepath)
    records = {}
    pages_to_ocr = []
    for page_num, page in enumerate(reader.pages, 1):
        try:
            text, needs_ocr = _pdf_native_page_text(page)
        except Exception as exc:
            # A parser/inspection failure must not certify a partial page as
            # complete. Try the existing bounded render/OCR path instead.
            logger.warning("PDF page %s needs OCR after text/image inspection failed: %s", page_num, exc)
            text, needs_ocr = "", True
        records[page_num] = {"page": page_num, "text": text, "method": "text"}
        if needs_ocr:
            pages_to_ocr.append(page_num)
    if pages_to_ocr:
        logger.info("PDF OCR required for %s of %s pages", len(pages_to_ocr), len(reader.pages))
        try:
            recovered = {int(item["page"]): item.get("text", "") for item in ocr_pdf_pages(filepath, pages_to_ocr)}
        except RuntimeError as exc:
            raise RuntimeError(
                f"Incomplete PDF extraction: image content on page(s) {', '.join(map(str, pages_to_ocr[:12]))} "
                f"requires OCR. {exc}"
            ) from exc
        for page_num in pages_to_ocr:
            native = records[page_num]["text"]
            ocr_text = recovered.get(page_num, "").strip()
            if not ocr_text:
                rendered = None
                try:
                    rendered = render_pdf_page_for_ocr(filepath, page_num, grayscale=False)
                    if _is_nearly_blank_page(rendered):
                        logger.info("Skipped nearly blank PDF page %s with no readable text", page_num)
                        continue
                    if _is_blank_ruled_page(rendered):
                        logger.info("Skipped blank ruled PDF page %s with no readable text", page_num)
                        continue
                    if _is_centered_color_graphic_page(rendered):
                        logger.info(
                            "Skipped image-only decorative PDF page %s with no readable text",
                            page_num,
                        )
                        continue
                except Exception as exc:
                    # Classification is optional. Uncertainty keeps the
                    # existing fail-closed behavior for unreadable content.
                    logger.debug("Could not inspect empty OCR result on page %s: %s", page_num, exc)
                finally:
                    close = getattr(rendered, "close", None)
                    if callable(close):
                        close()
                raise RuntimeError(
                    f"Incomplete PDF extraction: OCR returned no readable text on page {page_num}. "
                    "Review the page image quality before retrying."
                )
            from .extraction_quality import extraction_issues
            native_issues = extraction_issues(native)
            if native_issues and not extraction_issues(ocr_text):
                # Retaining both versions would keep contradictory OCR quantities
                # searchable. Only replace when a clean full-page pass was obtained.
                records[page_num]['text'] = ocr_text
                records[page_num]['method'] = 'ocr_recovered'
                continue
            records[page_num]["text"] = _merge_pdf_native_and_ocr(native, ocr_text)
            records[page_num]["method"] = "mixed" if native else "ocr"
    return [records[number] for number in sorted(records)], len(reader.pages)


def extract_from_pdf(
    filepath: str,
    source_name: str = None,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> List[Tuple[str, Dict]]:
    all_chunks = []
    filename   = source_name or Path(filepath).name
    page_records, page_count = _extract_pdf_page_records(filepath)
    logger.info("PDF pages: %s — %s", page_count, filename)

    document_parents = document_article_parent_chunks(page_records, filename, category)
    section_parents = document_section_parent_chunks(page_records, filename, category)
    continuation_articles: Dict[int, Dict] = {}
    for _, metadata in document_parents:
        start_page = int(metadata.get("page") or 0)
        end_page = int(metadata.get("page_end") or start_page)
        article_meta = {
            key: metadata[key]
            for key in ("article_number", "article_title", "citation")
            if metadata.get(key)
        }
        for continued_page in range(start_page + 1, end_page + 1):
            continuation_articles.setdefault(continued_page, article_meta)

    for item in sorted(page_records, key=lambda value: int(value.get("page") or 0)):
        page_number = int(item.get("page") or 0)
        page_chunks = chunk_text(
            item.get("text", ""),
            source=filename,
            page=page_number,
            category=category,
            include_article_parents=False,
            leading_article_metadata=continuation_articles.get(page_number),
        )
        all_chunks.extend(mark_extraction_method(page_chunks, item.get("method", "text")))

    if document_parents or section_parents:
        extraction_methods = {item.get("method", "text") for item in page_records}
        method = "mixed" if len(extraction_methods) > 1 else next(iter(extraction_methods), "text")
        all_chunks.extend(mark_extraction_method(document_parents + section_parents, method))

    chunk_count_before_deduplication = len(all_chunks)
    all_chunks = deduplicate_ocr_chunks(all_chunks)
    removed_duplicates = chunk_count_before_deduplication - len(all_chunks)
    if removed_duplicates:
        logger.info(
            "Removed %s duplicate OCR chunks from %s before embedding",
            removed_duplicates,
            filename,
        )

    logger.info(f"Extracted {len(all_chunks)} chunks from {filename}")
    return all_chunks


def extract_from_docx(
    filepath: str,
    source_name: str = None,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
) -> List[Tuple[str, Dict]]:
    filename   = source_name or Path(filepath).name
    doc        = DocxDocument(filepath)
    full_text  = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                full_text.append(row_text)

    chunks = chunk_text(
        "\n\n".join(full_text),
        source=filename,
        page=0,
        category=category,
    )
    chunks.extend(document_section_parent_chunks(
        [{'page': 0, 'text': '\n\n'.join(full_text)}], filename, category))
    logger.info(f"Extracted {len(chunks)} chunks from {filename}")
    return chunks


def _normalize_document_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in (text or "").splitlines()]
    normalized = []
    blank_seen = False
    for line in lines:
        if not line:
            if normalized and not blank_seen:
                normalized.append("")
            blank_seen = True
            continue
        normalized.append(line)
        blank_seen = False
    return "\n".join(normalized).strip()


def extract_original_document_text(filepath: str) -> Dict:
    path = Path(filepath)
    extension = path.suffix.lower()

    if extension == ".pdf":
        page_records, page_count = _extract_pdf_page_records(filepath)
        pages = [{"page": item["page"], "text": item["text"]} for item in page_records if item.get("text")]
        return {
            "format": "pdf",
            "page_count": page_count,
            "pages": pages,
            "text": "\n\n".join(
                f"Page {item['page']}\n{item['text']}" for item in pages
            ).strip(),
        }

    if extension == ".docx":
        doc = DocxDocument(filepath)
        parts = []
        for paragraph in doc.paragraphs:
            text = _normalize_document_text(paragraph.text)
            if text:
                parts.append(text)
        for table_index, table in enumerate(doc.tables, 1):
            rows = []
            for row in table.rows:
                row_text = " | ".join(
                    _normalize_document_text(cell.text)
                    for cell in row.cells
                    if _normalize_document_text(cell.text)
                )
                if row_text:
                    rows.append(row_text)
            if rows:
                parts.append(f"Table {table_index}\n" + "\n".join(rows))
        text = "\n\n".join(parts).strip()
        return {
            "format": "docx",
            "page_count": 0,
            "pages": [{"page": 0, "text": text}] if text else [],
            "text": text,
        }

    if extension in IMAGE_EXTENSIONS:
        ensure_ocr_available()
        try:
            with Image.open(filepath) as image:
                text = ocr_image_to_text(image, path.name)
        except RuntimeError:
            raise
        except Exception as exc:
            raise ValueError("Unable to open this image for OCR.") from exc
        return {
            "format": extension.lstrip("."),
            "page_count": 1,
            "pages": [{"page": 1, "text": text}] if text else [],
            "text": text,
        }

    if extension == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported. Convert the document to DOCX or PDF before uploading."
        )
    raise ValueError(f"Unsupported file type: {extension}")


def _prepare_document_for_ingestion(
    filepath: str,
    doc_db_id: int,
    source_name: str = None,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
    extra_metadata: Optional[Dict[str, str]] = None,
) -> Tuple[str, List[str], List[Dict[str, object]]]:
    """Synchronously extract and chunk one document.

    This function intentionally contains the CPU- and disk-heavy parser work so
    the async wrapper can run the whole operation in a worker thread.
    """

    path      = Path(filepath)
    doc_id    = f"doc_{doc_db_id}_{uuid.uuid4().hex[:8]}"
    extension = path.suffix.lower()
    category  = normalize_kb_category(category)

    if extension == ".pdf":
        chunks_with_meta = extract_from_pdf(filepath, source_name, category)
    elif extension in IMAGE_EXTENSIONS:
        chunks_with_meta = extract_from_image(filepath, source_name, category)
    elif extension in (".docx", ".doc"):
        if extension == ".doc":
            raise ValueError(
                "Legacy .doc files are not supported. Convert the document to DOCX or PDF before uploading."
            )
        chunks_with_meta = extract_from_docx(filepath, source_name, category)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    if not chunks_with_meta:
        logger.warning(f"No text extracted from {filepath}")
        return doc_id, [], []

    indexing_metadata = {
        **(extra_metadata or {}),
        "category": category,
        "doc_id": doc_id,
        "document_id": doc_id,
        "doc_db_id": str(doc_db_id),
    }
    if not indexing_metadata.get("document_year"):
        inferred_document_year = infer_document_year(
            "\n".join(text for text, _metadata in chunks_with_meta),
            {
                **indexing_metadata,
                "original_filename": source_name or path.name,
            },
        )
        if inferred_document_year is not None:
            indexing_metadata["document_year"] = inferred_document_year
    texts, metadatas = prepare_chunks_for_indexing(chunks_with_meta, indexing_metadata)
    return doc_id, texts, metadatas


async def ingest_document(
    filepath: str,
    doc_db_id: int,
    source_name: str = None,
    category: Optional[str] = DEFAULT_KB_CATEGORY,
    extra_metadata: Optional[Dict[str, str]] = None,
) -> int:
    """Extract, chunk, embed, and index a document without blocking FastAPI."""

    display_name = source_name or Path(filepath).name
    _start_ingestion_status(display_name)
    succeeded = False
    failure_message: Optional[str] = None
    try:
        doc_id, texts, metadatas = await asyncio.to_thread(
            _prepare_document_for_ingestion,
            filepath,
            doc_db_id,
            source_name,
            category,
            extra_metadata,
        )
        if not texts:
            failure_message = "No readable text was extracted from the document."
            return 0

        chunk_count = await add_documents(texts, metadatas, doc_id)
        succeeded = chunk_count > 0
        if not succeeded:
            failure_message = "No searchable chunks were created from the document."
        return chunk_count
    except Exception as exc:
        failure_message = str(exc)
        raise
    finally:
        _finish_ingestion_status(
            display_name,
            succeeded=succeeded,
            error=failure_message,
        )


def save_upload(file_bytes: bytes, filename: str) -> str:
    upload_dir = Path(settings.UPLOAD_DIR)
    upload_dir.mkdir(parents=True, exist_ok=True)
    safe_name   = re.sub(r'[^\w.\-]', '_', filename)
    unique_name = f"{uuid.uuid4().hex[:8]}_{safe_name}"
    filepath    = upload_dir / unique_name
    with open(filepath, "wb") as f:
        f.write(file_bytes)
    return str(filepath)
